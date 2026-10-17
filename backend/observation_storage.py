"""Bounded-memory observation uploads and disk-backed duplicate aggregation."""
import csv
import gzip
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import tempfile
from functools import lru_cache
from uuid import uuid4
import zipfile
from xml.etree.ElementTree import ParseError
from .observation_qc import normalize_qc, csv_delimiter

from fastapi import HTTPException
from starlette.requests import ClientDisconnect

MAX_UPLOAD_BYTES = 2 * 1024**3
MAX_EXCEL_BYTES = 20 * 1024**2
STORAGE_ROOT = Path(os.getenv("OBSERVATION_STORAGE_DIR", str(Path(__file__).resolve().parent.parent / "data" / "observation_uploads")))


async def receive_observation_file(request, filename):
    if Path(filename).suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
        raise HTTPException(status_code=422, detail="Upload CSV, XLSX, or XLSM")
    limit = MAX_UPLOAD_BYTES if Path(filename).suffix.lower() == ".csv" else MAX_EXCEL_BYTES
    limit_message = "CSV observation uploads must be at most 2 GB; Excel files must be at most 20 MB (save large workbooks as CSV)"
    try:
        declared = int(request.headers.get("content-length", "0"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid Content-Length") from exc
    if declared < 0: raise HTTPException(status_code=400, detail="Invalid Content-Length")
    if declared > limit: raise HTTPException(status_code=413, detail=limit_message)
    STORAGE_ROOT.mkdir(parents=True, exist_ok=True)
    if shutil.disk_usage(STORAGE_ROOT).free < max(512 * 1024**2, declared * 4):
        raise HTTPException(status_code=507, detail="Not enough server disk space for this observation file and its processing index")
    path = STORAGE_ROOT / (uuid4().hex + Path(filename).suffix.lower())
    size = 0; digest = hashlib.sha256(); disk_check = 64 * 1024**2
    try:
        # Disk I/O runs off the event loop; no request.body() or full-file byte buffer.
        from fastapi.concurrency import run_in_threadpool
        with path.open("xb") as stream:
            async for chunk in request.stream():
                size += len(chunk)
                if size > limit:
                    raise HTTPException(status_code=413, detail=limit_message)
                await run_in_threadpool(stream.write, chunk)
                digest.update(chunk)
                if size >= disk_check:
                    if shutil.disk_usage(STORAGE_ROOT).free < 128 * 1024**2:
                        raise HTTPException(status_code=507, detail="Server disk space is low; the partial upload was removed")
                    disk_check = size + 64 * 1024**2
        if not size: raise HTTPException(status_code=422, detail="Choose a non-empty observation file")
        if declared and size != declared: raise HTTPException(status_code=400, detail="Observation upload was incomplete")
        return path, dict(storage_key=path.name, size_bytes=size, sha256=digest.hexdigest())
    except ClientDisconnect as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=
            "Observation upload was interrupted before the file finished transferring. The partial file was removed. Keep the page open and retry on a stable connection.") from exc
    except BaseException:
        path.unlink(missing_ok=True)
        raise


class DiskObservations:
    def __init__(self, root):
        self.directory = tempfile.TemporaryDirectory(prefix="index-", dir=root)
        self.connection = sqlite3.connect(str(Path(self.directory.name) / "observations.sqlite"))
        # This index is disposable and never reused after failure; source/report storage is separate.
        self.connection.execute("PRAGMA journal_mode=OFF")
        self.connection.execute("PRAGMA synchronous=OFF")
        self.connection.execute("PRAGMA cache_size=-65536")
        self.connection.execute("PRAGMA temp_store=FILE")
        self.connection.execute("""CREATE TABLE observations (
            kind INTEGER, station TEXT, element TEXT, day TEXT, slot INTEGER, value REAL, qc_status TEXT,
            PRIMARY KEY(kind,station,element,day,slot)) WITHOUT ROWID""")

    def insert(self, rows):
        self.connection.executemany("""INSERT INTO observations VALUES (?,?,?,?,?,?,?)
            ON CONFLICT(kind,station,element,day,slot) DO UPDATE
            SET value=COALESCE(excluded.value,observations.value),
                qc_status=CASE WHEN excluded.value IS NOT NULL THEN excluded.qc_status ELSE observations.qc_status END""", rows)
        self.connection.commit()

    @lru_cache(maxsize=32)
    def day(self, kind, station, element, day):
        return dict(self.connection.execute("SELECT slot,value FROM observations WHERE kind=? AND station=? AND element=? AND day=?",
            (kind, station, element, day)))

    def close(self):
        self.day.cache_clear()
        self.qc_day.cache_clear()
        self.connection.close()
        self.directory.cleanup()

    @lru_cache(maxsize=32)
    def qc_day(self, kind, station, element, day):
        return dict(self.connection.execute("SELECT slot,qc_status FROM observations WHERE kind=? AND station=? AND element=? AND day=? AND value IS NOT NULL ORDER BY slot",
            (kind, station, element, day)))


class ObservationLookup:
    def __init__(self, store, kind): self.store, self.kind = store, kind
    def get(self, key, default=None):
        station, element, day, slot = key
        return self.store.day(self.kind, station, element, day.isoformat()).get(slot, default)


class DailyLookup:
    def __init__(self, store): self.store = store
    def __contains__(self, key):
        station, element, day = key
        return any(value is not None for value in self.store.day(0, station, element, day.isoformat()).values())


class QCLookup:
    def __init__(self, store, kind): self.store, self.kind = store, kind
    def get(self, key, default="none"):
        station, element, day, *slot = key
        values = self.store.qc_day(self.kind, station, element, day.isoformat())
        return values.get(slot[0], default) if slot else next(reversed(values.values()), default)


def parse_disk_source(path, filename):
    from .data_upload_engine import parse_datetime, normalized_code, present_value, observation_column_indexes
    store = DiskObservations(path.parent)
    workbook = stream = None
    try:
        if path.suffix.lower() == ".csv":
            stream = path.open("r", encoding="utf-8-sig", newline="", buffering=1024 * 1024)
            first = stream.readline(1024 * 1024); stream.seek(0)
            rows = iter(csv.reader(stream, delimiter=csv_delimiter(first)))
        else:
            from openpyxl import load_workbook
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024**2:
                    raise ValueError("Workbook expands beyond 100 MB; save it as CSV for large uploads")
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            rows = iter(workbook.active.iter_rows(values_only=True))
        columns = observation_column_indexes(next(rows, []))
        last_column = max(value for name, value in columns.items() if name != "qc_status")
        periods, stations, elements = set(), set(), set()
        valid = invalid = 0; batch = []
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row): continue
            if len(row) <= last_column: invalid += 1; continue
            code, element = normalized_code(row[columns["station_id"]]), normalized_code(row[columns["element_id"]])
            if not code or not element or len(code) > 50 or len(element) > 30:
                invalid += 1; continue
            try: stamp = parse_datetime(row[columns["date_time"]])
            except (ValueError, OverflowError, TypeError): invalid += 1; continue
            periods.add(stamp.strftime("%Y-%m")); stations.add(code); elements.add(element)
            if len(periods) > 12: raise ValueError("Upload at most 12 different months at a time")
            if len(stations) > 500 or len(elements) > 2000:
                raise ValueError("Upload at most 500 station codes and 2,000 element codes at a time")
            value = present_value(row[columns["value"]]); day = stamp.date().isoformat(); valid += 1
            try: status = normalize_qc(row[columns["qc_status"]] if "qc_status" in columns and columns["qc_status"] < len(row) else None)
            except ValueError as exc: raise ValueError(f"Row {row_number}: {exc}") from exc
            batch.extend([(0, code, element, day, stamp.hour * 60 + stamp.minute, value, status),
                          (1, code, element, day, stamp.hour, value, status)])
            if len(batch) >= 10000:
                store.insert(batch); batch.clear()
                if shutil.disk_usage(path.parent).free < 128 * 1024**2:
                    raise ValueError("Server disk space is low; processing was stopped safely")
        if batch: store.insert(batch)
        if not periods: raise ValueError("No valid dated observations were found")
        counts = dict(store.connection.execute("SELECT kind,COUNT(*) FROM observations GROUP BY kind"))
        return dict(observations=ObservationLookup(store, 1), minute_observations=ObservationLookup(store, 0),
            qc_observations=QCLookup(store, 1), qc_minute_observations=QCLookup(store, 0), qc_daily=QCLookup(store, 0),
            daily=DailyLookup(store), periods=sorted(periods), station_codes=stations, element_codes=elements,
            row_count=valid, invalid_rows=invalid, duplicate_rows=valid-counts.get(1, 0),
            minute_duplicate_rows=valid-counts.get(0, 0), close=store.close)
    except (csv.Error, zipfile.BadZipFile, ParseError, EOFError) as exc:
        store.close()
        raise ValueError("The observation file is malformed or contains an oversized CSV field") from exc
    except BaseException:
        store.close()
        raise
    finally:
        if workbook is not None: workbook.close()
        if stream is not None: stream.close()


def store_result_details(result):
    path = STORAGE_ROOT / (uuid4().hex + ".json.gz")
    try:
        with gzip.open(path, "wt", encoding="utf-8") as stream:
            json.dump(result, stream)
        summary = {key: value for key, value in result.items() if key not in {"daily_details", "hourly_details"}}
        summary["details_file"] = path.name
        return summary, path
    except BaseException:
        path.unlink(missing_ok=True)
        raise


def load_result_details(summary):
    key = summary.get("details_file")
    if not key: return summary
    if Path(key).name != key or not key.endswith(".json.gz"):
        raise HTTPException(status_code=500, detail="Invalid observation report storage reference")
    try:
        with gzip.open(STORAGE_ROOT / key, "rt", encoding="utf-8") as stream:
            return json.load(stream)
    except OSError as exc:
        raise HTTPException(status_code=503, detail="Observation report storage is unavailable; restore the observation storage backup") from exc
