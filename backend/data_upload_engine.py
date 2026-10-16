"""Deterministic daily/hourly completeness checks translated from Missing_DH VBA."""
import calendar
from copy import deepcopy
from pathlib import Path
import csv
import io
import math
import re
import zipfile
from xml.etree.ElementTree import ParseError
from datetime import date, datetime, timedelta
from html import escape
from pydantic import BaseModel, Field, field_validator, model_validator
from typing import Literal

UPLOAD_SCHEMA = (
    """CREATE TABLE IF NOT EXISTS data_upload_config (
        config_id INT PRIMARY KEY, version INT NOT NULL, configuration JSON NOT NULL,
        updated_by_username VARCHAR(50) NOT NULL, updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP)""",
    """CREATE TABLE IF NOT EXISTS data_upload_runs (
        run_id INT AUTO_INCREMENT PRIMARY KEY, filename VARCHAR(255) NOT NULL,
        source_data LONGBLOB NOT NULL, config_version INT NOT NULL, config_snapshot JSON NOT NULL,
        result_data LONGTEXT NOT NULL, pdf_data LONGBLOB NOT NULL,
        created_by_user_id INT NULL, created_by_username VARCHAR(50) NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""",
    """CREATE TABLE IF NOT EXISTS data_upload_results (
        result_id INT AUTO_INCREMENT PRIMARY KEY, run_id INT NOT NULL, station_id INT NOT NULL,
        station_code VARCHAR(50) NOT NULL, station_name VARCHAR(200) NOT NULL, district VARCHAR(100) NULL,
        report_month DATE NOT NULL, daily_expected INT NOT NULL, daily_available INT NOT NULL,
        daily_missing INT NOT NULL, hourly_expected INT NOT NULL, hourly_available INT NOT NULL,
        hourly_missing INT NOT NULL, daily_percent DECIMAL(6,2) NOT NULL,
        hourly_percent DECIMAL(6,2) NOT NULL, overall_percent DECIMAL(6,2) NOT NULL,
        UNIQUE KEY uq_upload_result (run_id, station_id, report_month),
        FOREIGN KEY (run_id) REFERENCES data_upload_runs(run_id),
        FOREIGN KEY (station_id) REFERENCES stations(station_id),
        INDEX idx_upload_latest (station_id, report_month, run_id))""",
)

DAILY = dict(zip(["2","3","5","18","85","477","573","574","575","576"],
    ["Tmax","Tmin","Rainfall","Daily bac","Grass min","Evaporation Piche Daily","Evaporation piche 06:00-18:00","Evaporation piche 18:00-06:00","Evaporation Bac 06:00-18:00","Evaporation Bac 18:00-06:00"]))
HOURLY = dict(zip(["112","196","110","167","851","190","193","195","115","191","192","194","170","201","171","104","101","102","105","166","103","183","184","106","185","532","176","177","179","180"],
    ["Wind Direction","Wind Speed","Visibility","Present Weather","Past Weather","Low Cloud Type","Middle Cloud Type","High Cloud Type","Total Cloud Cover","Low Cloud Amount","Low Cloud Height","Middle Cloud Amount","Middle Cloud Height","High Cloud Amount","High Cloud Height","Precipitation hourly","Dry Bulb Temperature","Wet Bulb Temperature","Relative Humidity","Vapor Pressure","Dew Point Temperature","Thermograph Temperature","Hygrograph Humidity","Station Pressure","Geopotential Height","Sunshine Duration","Soil Temperature at 10 cm","Soil Temperature at 20 cm","Soil Temperature at 50 cm","Soil Temperature at 100 cm"]))
STATIONS = dict(zip(["10101100","50410100","20515100","20602100","30109100","40302100","40303100","40503100","50210100","50604100","30304100"],
    ["GITEGA","KAWANGIRE","GIKONGORO","BYIMANA","RUBENGERA","BUSOGO","RUHENGERI","BYUMBA","NYAGATARE","KIBUNGO","GISENYI"]))
EXCLUDED = {"50410100":["177","575","576","18"],"20515100":["176"],"30109100":["177","575","576","18"],
    "40302100":["177"],"40303100":["177","85"],"40503100":["176","177","575","576","18"],
    "50210100":["177"],"50604100":["177"],"30304100":["176","179","190"]}
DAYLIGHT = set(["112","196","110","167","851","190","193","195","115","191","192","194","170","201","171","532"])
CLOUD = {"191":"190","192":"190","194":"193","170":"193","201":"195","171":"195"}


def observation_defaults():
    defaults = {}
    for code in "1022 1012 1081 2512 2132 1151 1032 1272 1132 1107 2122 2422 2222 2322 2112 2412 2212 2312 2181 1441 1432 1042 1092".split():
        defaults[code] = dict(cadence="hourly", schedule_mode="interval", interval_minutes=10,
                              times=[], hours=list(range(24)), enabled=True)
    hourly = "112 196 110 167 851 190 193 195 115 191 192 194 170 201 171 116 117 118 120 121 122 124 125 126 128 129 130 104 101 102 105 166 103 183 184 106 185 532"
    for code in hourly.split():
        defaults[code] = dict(cadence="hourly", schedule_mode="interval", interval_minutes=60,
            times=[], hours=list(range(6, 19)) if code in DAYLIGHT else list(range(24)), enabled=True)
    for codes, times in [("176 177 179 180", ["06:00", "12:00", "18:00"]),
                         ("2 3 5 85 477 18", ["08:00"]), ("573 575", ["18:00"]), ("574 576", ["06:00"])]:
        for code in codes.split():
            defaults[code] = dict(cadence="daily" if len(times) == 1 else "hourly", schedule_mode="times",
                interval_minutes=60, times=list(times), hours=[int(t[:2]) for t in times], enabled=True)
    return defaults


def apply_observation_defaults(configuration):
    """Explicit owner-approved one-time defaults; never run over subsequent admin edits."""
    config = deepcopy(configuration)
    previous = dict(config.get("archived_parameter_rules", {}))
    previous.update({p["element_id"]: p for p in config["parameters"]})
    for code, values in observation_defaults().items():
        rule = previous.setdefault(code, dict(element_id=code))
        rule.update(values)
        if code in CLOUD: rule.setdefault("zero_parent", CLOUD[code])
    config["parameters"] = [previous[p["element_id"]] for p in config["parameters"]]
    config["archived_parameter_rules"] = previous
    config["schedule_defaults_version"] = 1
    return config


class ParameterRule(BaseModel):
    element_id: str = Field(min_length=1, max_length=30, pattern=r"^[A-Za-z0-9_-]+$")
    name: str = Field(min_length=1, max_length=150)
    cadence: str = Field(pattern="^(daily|hourly)$")
    hours: list[int] = Field(default_factory=list, max_length=24)
    enabled: bool = True
    station_codes: list[str] = Field(default_factory=list, max_length=500)
    assessment_station_codes: list[str] | None = Field(default=None, max_length=500)
    excluded_station_codes: list[str] = Field(default_factory=list, max_length=500)
    zero_parent: str | None = None
    schedule_mode: str = Field(default="legacy", pattern="^(none|legacy|interval|times)$")
    interval_minutes: int = Field(default=60, ge=1, le=1440)
    times: list[str] = Field(default_factory=list, max_length=1440)

    @field_validator("times")
    @classmethod
    def valid_times(cls, values):
        if any(not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value) for value in values):
            raise ValueError("Observation times must use HH:MM (00:00 to 23:59)")
        if len(set(values)) != len(values):
            raise ValueError("Observation times must be unique")
        return sorted(values)

    @model_validator(mode="after")
    def valid_schedule(self):
        if self.schedule_mode == "none" and self.enabled:
            raise ValueError("Set an observation schedule before enabling this parameter")
        if self.schedule_mode == "times" and not self.times:
            raise ValueError("Select at least one observation time")
        if self.schedule_mode == "interval" and self.cadence != "hourly":
            raise ValueError("Interval schedules use the intraday/hourly counting group")
        if self.schedule_mode == "times" and self.cadence == "daily" and len(self.times) != 1:
            raise ValueError("Daily counting requires exactly one observation time")
        return self

    @field_validator("hours")
    @classmethod
    def valid_hours(cls, values):
        if len(set(values)) != len(values) or any(hour < 0 or hour > 23 for hour in values):
            raise ValueError("Hours must be different integers from 0 to 23")
        return sorted(values)

    @field_validator("name")
    @classmethod
    def name_not_blank(cls, value):
        if not value.strip(): raise ValueError("Enter the parameter name")
        return value.strip()


class HourOverride(BaseModel):
    hours: list[int] = Field(default_factory=list, max_length=24)
    reported_excluded_hours: list[int] = Field(default_factory=list, max_length=24)
    count_in_summary: bool = True
    @field_validator("hours", "reported_excluded_hours")
    @classmethod
    def valid_hours(cls, values):
        return ParameterRule.valid_hours(values)

    @model_validator(mode="after")
    def disjoint_hours(self):
        if set(self.hours) & set(self.reported_excluded_hours):
            raise ValueError("Expected and excluded hours cannot overlap")
        return self


class StationRule(BaseModel):
    station_code: str = Field(min_length=1, max_length=50, pattern=r"^[A-Za-z0-9_-]+$")
    station_name: str = Field(min_length=1, max_length=200)
    enabled: bool = True
    daily_excluded: list[str] = Field(default_factory=list, max_length=200)
    hourly_overrides: dict[str, HourOverride] = Field(default_factory=dict, max_length=200)


class UploadConfiguration(BaseModel):
    version: int = Field(ge=1)
    parameters: list[ParameterRule] = Field(min_length=1, max_length=200)
    stations: list[StationRule] = Field(min_length=1, max_length=500)
    station_statuses: list[Literal["Operational", "Under maintenance", "Suspended", "Closed"]] = Field(
        default_factory=lambda: ["Operational", "Under maintenance", "Suspended"], min_length=1, max_length=4)

    @model_validator(mode="after")
    def references(self):
        params = {item.element_id: item for item in self.parameters}
        if len(params) != len(self.parameters) or len({s.station_code for s in self.stations}) != len(self.stations):
            raise ValueError("Station and parameter codes must be unique")
        for parameter in self.parameters:
            if parameter.zero_parent and (parameter.zero_parent not in params or params[parameter.zero_parent].cadence != "hourly" or parameter.zero_parent == parameter.element_id):
                raise ValueError("A zero-cloud parent must be a different hourly parameter")
        for station in self.stations:
            if not set(station.daily_excluded).issubset(params): raise ValueError("An exclusion references an unknown parameter")
            if any(code not in params or params[code].cadence != "hourly" for code in station.hourly_overrides):
                raise ValueError("Hourly overrides must reference hourly parameters")
        return self


def default_configuration():
    parameters = [dict(element_id=code, name=name, cadence="daily", hours=[]) for code, name in DAILY.items()]
    for code, name in HOURLY.items():
        hours = [6,12,18] if code in {"176","177","179","180"} else list(range(6,19)) if code in DAYLIGHT else list(range(24))
        parameters.append(dict(element_id=code, name=name, cadence="hourly", hours=hours,
            station_codes=["30304100","40303100"] if code in {"106","185"} else [], zero_parent=CLOUD.get(code)))
    stations = [dict(station_code=code, station_name=name, daily_excluded=EXCLUDED.get(code, []), hourly_overrides={}) for code, name in STATIONS.items()]
    busogo = next(s for s in stations if s["station_code"] == "40302100")
    busogo["hourly_overrides"] = {code: dict(hours=list(range(6,19)), reported_excluded_hours=list(range(6))+list(range(19,24)), count_in_summary=True) for code in ["105","102","103","166"]}
    busogo["hourly_overrides"]["184"] = dict(hours=[], reported_excluded_hours=list(range(24)), count_in_summary=False)
    return UploadConfiguration(version=1, parameters=parameters, stations=stations).model_dump()


def normalized_code(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isfinite(value) and value == int(value): return str(int(value))
    text = str(value or "").strip()
    if re.fullmatch(r"\d+\.0+", text): return text.split(".")[0]
    return text


def parse_datetime(value):
    if isinstance(value, datetime): parsed = value
    elif isinstance(value, date): parsed = datetime.combine(value, datetime.min.time())
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel
        parsed = from_excel(value)
        if not isinstance(parsed, datetime): raise ValueError("Invalid Excel date")
    else:
        text = str(value or "").strip()
        if len(text) >= 10 and text[4:5] == "-":
            try:
                parsed = datetime.fromisoformat(text)
                if parsed.tzinfo is None and parsed.year >= 1900: return parsed
            except ValueError: pass
        parsed = None
        for fmt in ["%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S"]:
            try: parsed = datetime.strptime(text, fmt); break
            except ValueError: pass
        if parsed is None: raise ValueError("Use DD/MM/YYYY HH:MM or an Excel date")
    if parsed.year < 1900: raise ValueError("Invalid observation year")
    return parsed


def present_value(value):
    if isinstance(value, bool) or value is None or not str(value).strip(): return None
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError): return None


def observation_column_indexes(headers):
    normalized = [str(value or "").strip().lower() for value in headers]
    columns = {}
    for name in ["station_id", "element_id", "element_name", "date_time", "value"]:
        matches = [index for index, value in enumerate(normalized)
                   if value == name or (name == "date_time" and value.startswith("date_time"))]
        if len(matches) != 1:
            raise ValueError("Expected columns: station_id, element_id, element_name, date_time, value (once each, in any order; extra columns are allowed)")
        columns[name] = matches[0]
    return columns


def parse_source(content, filename):
    if isinstance(content, Path):
        from .observation_storage import parse_disk_source
        return parse_disk_source(content, filename)
    extension = filename.rsplit(".", 1)[-1].lower()
    workbook = None
    if extension == "csv":
        try: text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc: raise ValueError("CSV must use UTF-8 encoding") from exc
        delimiter = ";" if text.splitlines() and text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
        rows = iter(csv.reader(io.StringIO(text), delimiter=delimiter))
    elif extension in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond 100 MB; split it into smaller uploads")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            rows = iter(workbook.active.iter_rows(values_only=True))
        except ValueError: raise
        except Exception as exc: raise ValueError("Invalid Excel workbook") from exc
    else: raise ValueError("Upload CSV, XLSX, or XLSM; legacy XLS must be saved as XLSX")
    try:
        columns = observation_column_indexes(next(rows, []))
        last_column = max(columns.values())
        obs, minute_obs, daily, periods, codes = {}, {}, set(), set(), set()
        invalid = duplicates = minute_duplicates = valid = 0
        for index, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row): continue
            if len(row) <= last_column: invalid += 1; continue
            code, element = normalized_code(row[columns["station_id"]]), normalized_code(row[columns["element_id"]])
            try: dt = parse_datetime(row[columns["date_time"]])
            except (ValueError, OverflowError, TypeError): invalid += 1; continue
            if not code or not element: invalid += 1; continue
            period = dt.strftime("%Y-%m"); periods.add(period); codes.add(code); valid += 1
            key = (code, element, dt.date(), dt.hour)
            if key in obs: duplicates += 1
            value = present_value(row[columns["value"]])
            minute_key = (code, element, dt.date(), dt.hour * 60 + dt.minute)
            if minute_key in minute_obs: minute_duplicates += 1
            if value is not None or minute_key not in minute_obs:
                minute_obs[minute_key] = value
            if value is not None:
                obs[key] = value
                daily.add((code, element, dt.date()))
            elif key not in obs: obs[key] = None
        if not periods: raise ValueError("No valid dated observations were found")
        if len(periods) > 12: raise ValueError("Upload at most 12 different months at a time")
        return dict(observations=obs, minute_observations=minute_obs, daily=daily, periods=sorted(periods), station_codes=codes, row_count=valid,
            invalid_rows=invalid, duplicate_rows=duplicates, minute_duplicate_rows=minute_duplicates)
    except (csv.Error, ParseError, zipfile.BadZipFile, EOFError) as exc:
        raise ValueError("The observation file is malformed or contains an oversized field") from exc
    finally:
        if workbook is not None: workbook.close()


def hour_ranges(hours):
    ranges = []
    for hour in sorted(set(hours)):
        if ranges and ranges[-1][1] + 1 == hour: ranges[-1][1] = hour
        else: ranges.append([hour, hour])
    return ", ".join(f"{start:02d}:00" + (f"-{end:02d}:00" if end != start else "") for start, end in ranges)


def scheduled_minutes(param, hours=None):
    if param.schedule_mode == "times":
        slots = [int(value[:2]) * 60 + int(value[3:]) for value in param.times]
        if hours is None: return slots
    else:
        slots = range(0, 1440, param.interval_minutes)
    allowed = set(param.hours if hours is None else hours)
    return [minute for minute in slots if param.cadence == "daily" or minute // 60 in allowed]


def time_ranges(minutes, step=1):
    ranges = []
    for minute in sorted(set(minutes)):
        if ranges and ranges[-1][1] + step == minute: ranges[-1][1] = minute
        else: ranges.append([minute, minute])
    def label(value): return f"{value // 60:02d}:{value % 60:02d}"
    return ", ".join(label(start) + ("-" + label(end) if end != start else "") for start, end in ranges)


def analyze_source(parsed, configuration, inventory=None):
    config = UploadConfiguration.model_validate(configuration)
    stations = [s for s in config.stations if s.enabled]
    if not stations: raise ValueError("No enabled stations are configured for this account")
    if configuration.get("station_scope") != "registry" and not parsed["station_codes"].intersection(s.station_code for s in stations):
        raise ValueError("No uploaded station codes match enabled station profiles. Use registered station codes or ask Admin to configure the stations.")
    if len(stations) * len(parsed["periods"]) > 6000: raise ValueError("Limit the upload to 6,000 station-month assessments")
    params = {p.element_id: p for p in config.parameters}
    mapped = set(inventory.get("mapped_elements", [])) if inventory is not None else set()
    assignments = inventory.get("assignments", []) if inventory is not None else []
    availability = {}
    for assignment in assignments:
        for element in assignment["element_ids"]:
            availability.setdefault((assignment["station_code"], element), []).append(assignment)

    def operational(code, element, current):
        if element not in mapped:
            return True
        if inventory.get("availability_mode") == "current_status":
            return any(assignment["status"] == "Operational"
                       for assignment in availability.get((code, element), []))
        for assignment in availability.get((code, element), []):
            status = None
            for period in assignment["periods"]:
                if period["effective_date"] <= current.isoformat():
                    status = period["status"]
                else:
                    break
            if status == "Operational":
                return True
        return False
    def downtime_status(code, element, current):
        statuses = []
        for assignment in availability.get((code, element), []):
            status = assignment["status"] if inventory.get("availability_mode") == "current_status" else "Not installed / no status recorded"
            if inventory.get("availability_mode") != "current_status":
                for period in assignment.get("periods", []):
                    if period["effective_date"] <= current.isoformat(): status = period["status"]
                    else: break
            if status != "Operational" and status not in statuses: statuses.append(status)
        return "; ".join(statuses) or "Not installed / no status recorded"

    def eligible(param, code):
        if code in param.excluded_station_codes: return False
        if configuration.get("catalog_managed"):
            selected = param.assessment_station_codes if param.assessment_station_codes is not None else param.station_codes
            return code in selected and code in inventory.get("eligible_stations", {}).get(param.element_id, [])
        return not param.station_codes or code in param.station_codes
    summaries, daily_rows, hourly_rows, element_rows = [], [], [], []
    obs, daily = parsed["observations"], parsed["daily"]
    minute_obs = parsed.get("minute_observations", {})
    work = 0
    for station in stations:
        code = station.station_code
        for period in parsed["periods"]:
            year, month = map(int, period.split("-")); days = calendar.monthrange(year, month)[1]
            counts = dict(daily_expected=0, daily_available=0, hourly_expected=0, hourly_available=0)
            element_counts = {}
            def credit(param, available=False):
                values = element_counts.setdefault(param.element_id, [0, 0])
                values[0] += 1
                values[1] += int(available)
            daily_params = [p for p in config.parameters if p.enabled and p.schedule_mode != "none" and p.cadence == "daily" and eligible(p, code)
                and (p.element_id in mapped or p.element_id not in station.daily_excluded)]
            excluded = [params[e].name for e in station.daily_excluded if e not in mapped]
            for day in range(1, days+1):
                current = date(year, month, day); missing = []; day_excluded = list(excluded); statuses = []; downtime = []
                for param in daily_params:
                    element_counts.setdefault(param.element_id, [0, 0])
                    if not operational(code, param.element_id, current):
                        status = downtime_status(code, param.element_id, current)
                        day_excluded.append(param.name + " (" + status + ")")
                        statuses.append(param.name + ": " + status)
                        downtime.append(dict(element_name=param.name, status=status))
                        continue
                    counts["daily_expected"] += 1
                    available = ((code, param.element_id, current) in daily if param.schedule_mode == "legacy"
                        else minute_obs.get((code, param.element_id, current, scheduled_minutes(param)[0])) is not None)
                    credit(param, available)
                    if available: counts["daily_available"] += 1
                    else: missing.append(param.name)
                if missing or day_excluded: daily_rows.append(dict(station_code=code, station_name=station.station_name, date=current.isoformat(), missing=missing, excluded=day_excluded, entered="; ".join(statuses) or "No", downtime=downtime))
            for param in config.parameters:
                if not param.enabled or param.schedule_mode == "none" or param.cadence != "hourly" or not eligible(param, code): continue
                override = station.hourly_overrides.get(param.element_id)
                element_counts.setdefault(param.element_id, [0, 0])
                hours = override.hours if override else param.hours
                marked = override.reported_excluded_hours if override else []
                count = override.count_in_summary if override else True
                exact = param.schedule_mode != "legacy"
                slots = scheduled_minutes(param, override.hours if override else None) if exact else hours
                work += days * len(slots)
                if work > 100000000:
                    raise ValueError("This schedule exceeds 100 million checks. Upload fewer months or enable fewer stations/elements.")
                for day in range(1, days+1):
                    current = date(year, month, day); missing = []
                    if not operational(code, param.element_id, current):
                        if hours or marked:
                            hourly_rows.append(dict(station_code=code, station_name=station.station_name,
                                element_id=param.element_id, element_name=param.name,
                                date=current.isoformat(), missing_hours="", excluded_hours=hour_ranges(set(hours) | set(marked)),
                                entered=downtime_status(code, param.element_id, current)))
                        continue
                    for hour in slots:
                        observations = minute_obs if exact else obs
                        value = observations.get((code, param.element_id, current, hour))
                        if value is None and param.zero_parent and operational(code, param.zero_parent, current) and observations.get((code, param.zero_parent, current, hour)) == 0:
                            continue
                        if count:
                            counts["hourly_expected"] += 1
                            credit(param, value is not None)
                            if value is not None: counts["hourly_available"] += 1
                        if value is None: missing.append(hour)
                    if missing or marked:
                        hourly_rows.append(dict(station_code=code, station_name=station.station_name, element_id=param.element_id,
                            element_name=param.name, date=current.isoformat(), missing_hours=(time_ranges(missing, param.interval_minutes if param.schedule_mode == "interval" else 1) if exact else hour_ranges(missing)), excluded_hours=hour_ranges(marked)))
            for cadence in ["daily", "hourly"]:
                expected, available = counts[f"{cadence}_expected"], counts[f"{cadence}_available"]
                counts[f"{cadence}_missing"] = expected - available
                counts[f"{cadence}_percent"] = round(100 * available / expected, 2) if expected else 0
            expected = counts["daily_expected"] + counts["hourly_expected"]
            available = counts["daily_available"] + counts["hourly_available"]
            summaries.append(dict(station_code=code, station_name=station.station_name, report_month=period, **counts,
                overall_percent=round(100 * available / expected, 2) if expected else None))
            for element_id, (target, received) in element_counts.items():
                if configuration.get("catalog_managed") and code not in inventory.get("eligible_stations", {}).get(element_id, []):
                    continue
                element_rows.append(dict(station_code=code, station_name=station.station_name, report_month=period,
                    element_id=element_id, element_name=params[element_id].name, cadence=params[element_id].cadence, expected=target, available=received,
                    missing=target-received, score=round(100*received/target, 2) if target else None))
    unknown = sorted(parsed["station_codes"] - {s.station_code for s in stations})
    warnings = []
    if inventory is not None:
        warnings.append("Expected counts use the instrument availability snapshot saved with this upload; only Operational instruments count. Unmapped elements retain the configured rules.")
        if inventory.get("availability_mode") == "current_status":
            warnings.append("Current instrument status applies to the entire uploaded period, including past months; historical outage dates are not used.")
        else:
            warnings.append("Dated station-instrument status history is used. Downtime days are excluded even when the instrument is currently Operational. Status changes take effect at 00:00 on their effective date; missing historical dates cannot be inferred.")
        unconfigured_elements = sorted(mapped - set(params))
        if unconfigured_elements:
            warnings.append("Catalog element IDs without a parameter schedule are not assessed: " + ", ".join(unconfigured_elements))
        unmapped = sorted({p.element_id for p in config.parameters if p.enabled} - mapped)
        if unmapped:
            warnings.append("No catalog instrument mapping; configured expectations retained for: " + ", ".join(unmapped))
    if parsed["invalid_rows"]: warnings.append(f'{parsed["invalid_rows"]} invalid rows were skipped.')
    exact_schedules = any(p.schedule_mode != "legacy" for p in config.parameters)
    duplicate_count = parsed.get("minute_duplicate_rows", 0) if exact_schedules else parsed["duplicate_rows"]
    if duplicate_count: warnings.append(f'{duplicate_count} duplicate {"minute" if exact_schedules else "hourly"} keys: a present numeric value wins; the last present value controls cloud exemptions.')
    if exact_schedules:
        warnings.append("Scheduled observations are matched to the exact minute. Interval slots start at 00:00 each day; observations outside scheduled slots are not credited. Intraday/hourly totals include sub-hourly and fixed-time observations.")
    if unknown: warnings.append("Unconfigured/disabled station codes ignored: " + ", ".join(unknown))
    source_elements = parsed["element_codes"] if "element_codes" in parsed else {key[1] for key in obs}
    unknown_parameters = sorted(source_elements - set(params))
    if unknown_parameters: warnings.append("Unconfigured parameter codes ignored: " + ", ".join(unknown_parameters))
    unscheduled = sorted(source_elements & {p.element_id for p in config.parameters if not p.enabled or p.schedule_mode == "none"})
    if unscheduled: warnings.append("Disabled or unscheduled elements ignored: " + ", ".join(unscheduled))
    if any(set(s.daily_excluded) - mapped - {p.element_id for p in config.parameters if p.cadence == "daily"} for s in stations):
        warnings.append("Some original exclusions refer to hourly parameters. They remain annotations and do not reduce daily expected counts.")
    return dict(results=summaries, element_results=element_rows, daily_details=daily_rows, hourly_details=hourly_rows, periods=parsed["periods"],
        row_count=parsed["row_count"], invalid_rows=parsed["invalid_rows"], duplicate_rows=duplicate_count, warnings=warnings)


SUMMARY_COLUMNS = ["station_code","station_name","report_month","daily_expected","daily_available","daily_missing","hourly_expected","hourly_available","hourly_missing","daily_percent","hourly_percent","overall_percent"]
ELEMENT_COLUMNS = ["station_code", "station_name", "report_month", "element_id", "element_name", "expected", "available", "missing", "score"]


def csv_bytes(rows, columns):
    stream = io.StringIO(); writer = csv.writer(stream); writer.writerow(columns)
    for row in rows:
        values = [", ".join(row.get(key, [])) if isinstance(row.get(key), list) else row.get(key, "") for key in columns]
        writer.writerow(["'"+value if isinstance(value, str) and value.startswith(("=","+","-","@")) else value for value in values])
    return stream.getvalue().encode("utf-8-sig")


def downtime_summary(result):
    groups = {}
    for row in result.get("hourly_details", []):
        if row.get("entered"):
            groups.setdefault((row["station_code"], row["element_name"], row["entered"]), set()).add(row["date"])
    for row in result.get("daily_details", []):
        for entry in row.get("downtime", []):
            groups.setdefault((row["station_code"], entry["element_name"], entry["status"]), set()).add(row["date"])
    rows = []
    for (station, element, status), dates in sorted(groups.items()):
        ranges = []
        for value in sorted(dates):
            day = date.fromisoformat(value)
            if ranges and ranges[-1][1] + timedelta(days=1) == day: ranges[-1][1] = day
            else: ranges.append([day, day])
        label = ", ".join(str(start) + (" to " + str(end) if start != end else "") for start, end in ranges)
        rows.append([station, element, label, status])
    return rows


def report_pdf(result):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    stream = io.BytesIO(); styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    story = [Paragraph("Station data availability report", styles["Title"]), Paragraph(escape(", ".join(result["periods"])), styles["Normal"]), Spacer(1,12),
        Paragraph("Availability equals numeric observations received divided by observations expected under the saved rules. Zero is valid data. Excluded parameters and hours are not missing observations. Detailed daily and hourly exceptions are in the downloadable Word and CSV reports.", styles["Normal"]), Spacer(1,12)]
    if result.get("config_version"):
        story.extend([Paragraph(escape(f'Rule version {result["config_version"]}   Source {result.get("source_filename", "")}'), styles["Normal"]), Spacer(1,12)])
    for warning in result["warnings"]: story.extend([Paragraph(escape(warning), styles["Normal"]), Spacer(1,6)])
    headers = ["Station", "Month", "Daily exp.", "Daily avail.", "Daily miss.", "Hourly exp.", "Hourly avail.", "Hourly miss.", "Daily %", "Hourly %", "Overall %"]
    rows = [headers]
    for row in sorted(result["results"], key=lambda r: (r["report_month"], -(r["overall_percent"] if r["overall_percent"] is not None else -1), r["station_name"])):
        rows.append([Paragraph(escape(row["station_name"]+" "+row["station_code"]), styles["Normal"]), row["report_month"], *[str(row[key]) if row[key] is not None else "Not assessed" for key in SUMMARY_COLUMNS[3:]]])
    table = LongTable(rows, repeatRows=1, colWidths=[143,61]+[65]*9)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#dcece9")),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#b8ccca")),("FONTSIZE",(0,0),(-1,-1),8), ("VALIGN",(0,0),(-1,-1),"TOP"), ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f4f8f7")])]))
    story.append(table)
    if result.get("element_results"):
        story.extend([Spacer(1, 16), Paragraph("Scores by element", styles["Heading2"])])
        rows = [["Station", "Month", "Element ID", "Element name", "Expected", "Available", "Missing", "Score %"]]
        for item in result["element_results"]:
            rows.append([Paragraph(escape(str(item[key] if item[key] is not None else "Not assessed")), styles["Normal"])
                         for key in ["station_code", "report_month", "element_id", "element_name", "expected", "available", "missing", "score"]])
        element_table = LongTable(rows, repeatRows=1, colWidths=[80,65,65,205,70,70,70,65])
        element_table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey), ("VALIGN",(0,0),(-1,-1),"TOP")]))
        story.append(element_table)
    downtime = downtime_summary(result)
    if downtime:
        story.extend([Spacer(1, 16), Paragraph("Excluded instrument downtime", styles["Heading2"])])
        rows = [["Station", "Element", "Excluded dates", "Entered"]]
        rows.extend([[Paragraph(escape(value), styles["Normal"]) for value in row] for row in downtime])
        downtime_table = LongTable(rows, repeatRows=1, colWidths=[100,230,260,160])
        downtime_table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey), ("VALIGN",(0,0),(-1,-1),"TOP")]))
        story.append(downtime_table)
    doc.build(story); return stream.getvalue()


def report_details_zip(result):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("station_summary.csv", csv_bytes(result["results"], SUMMARY_COLUMNS))
        archive.writestr("element_scores.csv", csv_bytes(result.get("element_results", []), ELEMENT_COLUMNS))
        archive.writestr("daily_missing.csv", csv_bytes(result["daily_details"], ["station_code","station_name","date","missing","excluded","entered"]))
        archive.writestr("hourly_missing.csv", csv_bytes(result["hourly_details"], ["station_code","station_name","element_id","element_name","date","missing_hours","excluded_hours","entered"]))
    return stream.getvalue()


def report_docx(result):
    if len(result.get("element_results", [])) > 20000:
        raise ValueError("This Word report exceeds 20,000 station-element scores. Download CSV or Details ZIP, or select fewer stations or months.")
    condensed = len(result["daily_details"]) + len(result["hourly_details"]) + len(result.get("element_results", [])) > 20000
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    document = Document(); section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.65)
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Title"].font.color.rgb = RGBColor(0,0,0)
    def repeat_header(table):
        properties = table.rows[0]._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:tblHeader"))
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs: run.bold = True
    document.add_paragraph("Station data availability report", "Title")
    document.add_paragraph(", ".join(result["periods"]))
    if result.get("config_version"): document.add_paragraph(f'Rule version {result["config_version"]}   Source {result.get("source_filename", "")}')
    document.add_paragraph("Daily and hourly completeness under the configuration saved with this upload. Numeric zero is valid. Excluded entries are red and are not counted as missing. Full calendar months are assessed; this is not a range or plausibility quality-control test.")
    for warning in result["warnings"]: document.add_paragraph(warning)
    document.add_heading("Overall summary by station", level=1)
    summary_table = document.add_table(rows=1, cols=8); summary_table.style = "Table Grid"
    for cell, heading in zip(summary_table.rows[0].cells, ["Station", "Month", "Daily expected", "Daily available", "Hourly expected", "Hourly available", "Daily %", "Overall %"]): cell.text = heading
    for item in result["results"]:
        values = [item["station_name"], item["report_month"], item["daily_expected"], item["daily_available"], item["hourly_expected"], item["hourly_available"], item["daily_percent"], item["overall_percent"]]
        for cell, value in zip(summary_table.add_row().cells, values): cell.text = str(value) if value is not None else "Not assessed"
    repeat_header(summary_table)
    if result.get("element_results"):
        document.add_heading("Scores by element", level=1)
        table = document.add_table(rows=1, cols=len(ELEMENT_COLUMNS)); table.style = "Table Grid"
        for cell, key in zip(table.rows[0].cells, ELEMENT_COLUMNS): cell.text = key.replace("_", " ").title()
        for item in result["element_results"]:
            for cell, key in zip(table.add_row().cells, ELEMENT_COLUMNS): cell.text = str(item[key] if item[key] is not None else "Not assessed")
        repeat_header(table)
    if condensed:
        document.add_heading("Summary report for all selected stations", level=1)
        document.add_paragraph("All selected station and element scores are included above. Daily/hourly exceptions exceed 20,000 rows; download Details ZIP for the full exceptions. Recorded instrument downtime is grouped below.")
        downtime = downtime_summary(result)
        if downtime:
            document.add_heading("Excluded instrument downtime", level=2)
            table = document.add_table(rows=1, cols=4); table.style = "Table Grid"
            for cell, heading in zip(table.rows[0].cells,["Station","Element","Excluded dates","Entered"]): cell.text = heading
            for row in downtime:
                for cell, value in zip(table.add_row().cells,row): cell.text = value
            repeat_header(table)
        stream = io.BytesIO(); document.save(stream); return stream.getvalue()
    for row in result["results"]:
        code, month = row["station_code"], row["report_month"]
        document.add_heading(f'{row["station_name"]} {code} {month}', level=1)
        overall = str(row["overall_percent"]) + "%" if row["overall_percent"] is not None else "Not assessed"
        document.add_paragraph(f'Daily available {row["daily_available"]} of {row["daily_expected"]} ({row["daily_percent"]}%), missing {row["daily_missing"]}. Hourly available {row["hourly_available"]} of {row["hourly_expected"]} ({row["hourly_percent"]}%), missing {row["hourly_missing"]}. Overall available {overall}.')
        for heading, records, columns in [("Daily missing data", [d for d in result["daily_details"] if d["station_code"] == code and d["date"].startswith(month)], ["Date","Missing parameters","Excluded parameters","Entered"]),
            ("Hourly missing data", [d for d in result["hourly_details"] if d["station_code"] == code and d["date"].startswith(month)], ["Date","Parameter","Missing hours","Excluded hours","Entered"])]:
            document.add_heading(heading, level=2)
            if not records: document.add_paragraph("No missing observations or exclusions."); continue
            table = document.add_table(rows=1, cols=len(columns)); table.style = "Table Grid"
            for cell, name in zip(table.rows[0].cells, columns): cell.text = name
            repeat_header(table)
            for detail in records:
                values = [detail["date"], ", ".join(detail["missing"]), ", ".join(detail["excluded"]), detail.get("entered", "No")] if "missing" in detail else [detail["date"], detail["element_name"], detail["missing_hours"], detail["excluded_hours"], detail.get("entered", "No")]
                cells = table.add_row().cells
                for index, (cell, value) in enumerate(zip(cells, values)):
                    cell.text = str(value)
                    if index == len(values)-2:
                        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(190,35,45)
    stream = io.BytesIO(); document.save(stream); return stream.getvalue()
